#!/usr/bin/env python3
"""Convert a purchases .xlsx into a CSV ready for Lucrum's "Upload CSV" importer.

It reads your buying spreadsheet, maps its columns to the names the app expects
(regardless of small differences in spacing/case), drops the computed columns
(Total Cost and Landed Unit Cost -- the app recomputes those), formats dates as
YYYY-MM-DD, and writes a clean CSV.

Usage:
    python xlsx_to_purchases_csv.py purchases.xlsx
    python xlsx_to_purchases_csv.py purchases.xlsx -o out.csv --sheet "Sheet1"

Requires openpyxl:
    pip install openpyxl
"""

import argparse
import csv
import datetime as dt
import os
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed. Run:  pip install openpyxl")

# Canonical output columns the importer understands, in order.
OUTPUT_COLUMNS = [
    "Purchase Date", "Product Name", "My SKU", "Supplier / Source",
    "Product Link", "Order #", "Qty", "Unit Cost", "Sales Tax",
    "Shipping", "Status", "Notes",
]

# Normalized source header (lowercased, letters+digits only) -> output column.
# "Total Cost" / "Landed Unit Cost" are intentionally absent, so they're dropped.
HEADER_MAP = {
    "purchasedate": "Purchase Date", "date": "Purchase Date",
    "productname": "Product Name", "product": "Product Name",
    "mysku": "My SKU", "sku": "My SKU",
    "suppliersource": "Supplier / Source", "supplier": "Supplier / Source", "source": "Supplier / Source",
    "productlink": "Product Link", "link": "Product Link",
    "order": "Order #", "ordernumber": "Order #", "orderno": "Order #",
    "qty": "Qty", "quantity": "Qty",
    "unitcost": "Unit Cost", "cost": "Unit Cost",
    "salestax": "Sales Tax", "tax": "Sales Tax",
    "shipping": "Shipping",
    "status": "Status",
    "notes": "Notes",
}

DROP_COLUMNS = {"totalcost", "landedunitcost"}
REQUIRED = ("My SKU", "Qty", "Unit Cost")


def norm(header):
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def fmt(value):
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def main():
    ap = argparse.ArgumentParser(description="Convert a purchases .xlsx to a Lucrum-ready CSV.")
    ap.add_argument("xlsx", help="Path to the .xlsx file")
    ap.add_argument("-o", "--output", help="Output CSV path (default: <input>.csv)")
    ap.add_argument("--sheet", help="Worksheet name (default: first sheet)")
    ap.add_argument(
        "--sku-pad", type=int, default=14,
        help="Zero-pad purely numeric SKUs to this width. Excel silently strips the "
             "leading zeros from UPC-style SKUs like 00754127508214, which would stop "
             "them matching the catalog. Use 0 to disable. Default: 14",
    )
    args = ap.parse_args()

    if not os.path.exists(args.xlsx):
        sys.exit(f"File not found: {args.xlsx}")

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)

    # Find the header row: first row that maps at least two known columns.
    header, col_targets = None, None
    for row in rows:
        mapped = [HEADER_MAP.get(norm(c)) for c in row]
        if sum(1 for m in mapped if m) >= 2:
            header, col_targets = row, mapped
            break
    if header is None:
        sys.exit("Couldn't find a header row. Expected columns like 'My SKU', 'Qty', 'Unit Cost'.")

    idx_to_out = {i: t for i, t in enumerate(col_targets) if t}
    present = [c for c in OUTPUT_COLUMNS if c in idx_to_out.values()]
    dropped = [str(c) for c in header if norm(c) in DROP_COLUMNS]
    missing = [c for c in REQUIRED if c not in present]

    out_path = args.output or (os.path.splitext(args.xlsx)[0] + ".csv")
    written = 0
    padded = 0
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(present)
        for row in rows:  # generator continues right after the header row
            record = {idx_to_out[i]: fmt(v) for i, v in enumerate(row) if i in idx_to_out}
            if not any(v.strip() for v in record.values()):
                continue  # blank row
            sku = record.get("My SKU", "").strip()
            if not sku:
                continue  # importer requires a SKU
            # Excel stores UPC-style SKUs as numbers and drops the leading zeros.
            if args.sku_pad and sku.isdigit() and len(sku) < args.sku_pad:
                record["My SKU"] = sku.zfill(args.sku_pad)
                padded += 1
            writer.writerow([record.get(c, "") for c in present])
            written += 1

    print(f"Wrote {written} row(s) -> {out_path}")
    print(f"Columns: {', '.join(present)}")
    if padded:
        print(f"Restored leading zeros on {padded} numeric SKU(s) (Excel strips them) -> {args.sku_pad} digits")
    if dropped:
        print(f"Dropped computed columns (the app recomputes these): {', '.join(dropped)}")
    if missing:
        print(f"WARNING: required columns not found in the sheet: {', '.join(missing)}")


if __name__ == "__main__":
    main()
