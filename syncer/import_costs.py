"""Import exact per-line buying costs from the seller's SWIFTZAR workbooks.

Two books, two layouts, same meaning -- what was actually paid for each sale:

    SWIFTZAR LLC till Aug 8.xlsx   "Unit Purchase Price"  keyed by Customer Order ID
    Swiftzar Llc.xlsx              "Per Item Cost"        keyed by Order#

A spreadsheet row that names a specific customer order and SKU is the cost of
that exact sale, which beats any averaging. Where a line has no row of its own,
the SKU's other rows are used -- but only counted as known if they all agree.

Usage:
    python import_costs.py --dry-run
    python import_costs.py
"""

import argparse
import collections
import os
import re
import statistics

import openpyxl
from dotenv import load_dotenv
from supabase import create_client

load_dotenv()

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

# Two books, two layouts, same meaning. `charges` maps our column name to the
# header used in that book; a missing key just means that book does not record
# it (the Jul/Aug book predates third-party labels).
BOOKS = [
    {
        "file": "SWIFTZAR LLC till Aug 8.xlsx",
        "sheets": ["July 1- July 4", "July 4- Aug 8 ", "Aug 8-Aug 22"],
        "order": "customer order id",
        "sku": "sku/upc",
        "cost": "unit purchase price",
        "charges": {
            "walmart_fee": "walmart fees",
            "label_walmart": "wfs/ label fees",
            "warehouse_cost": "w.house cost",
        },
    },
    {
        "file": "Swiftzar Llc.xlsx",
        "sheets": ["Aug-26", "Sep-26", "Oct-26"],
        "order": "order#",
        "sku": "sku code",
        "cost": "per item cost",
        "charges": {
            "walmart_fee": "walmart fee",
            "label_walmart": "buy label from walmart",
            "label_third_party": "buy label from 3rd party",
            "warehouse_cost": "warehouse",
        },
    },
]

CHARGE_COLUMNS = ("walmart_fee", "label_walmart", "label_third_party", "warehouse_cost")


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _order_id(value) -> str:
    """Excel stores these as floats, so 200015460384635 arrives as ...635.0."""
    text = str(value or "").strip()
    return text[:-2] if text.endswith(".0") else text


def _number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def sku_variants(value):
    """A UPC typed into Excel loses its leading zeros and gains '.0'.

    '754127508214.0' in the sheet is '00754127508214' in Walmart's data.
    """
    text = str(value or "").strip()
    out = {text, text.lower()}
    if text.endswith(".0"):
        text = text[:-2]
        out |= {text, text.lower()}
    if text.isdigit():
        out |= {text.zfill(14), text.zfill(12)}
    return out


def read_books():
    rows = []
    for spec in BOOKS:
        filename = spec["file"]
        path = os.path.join(REPO, filename)
        if not os.path.exists(path):
            print(f"  skipping (not found): {filename}")
            continue
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in spec["sheets"]:
            if sheet not in book.sheetnames:
                continue
            ws, header, count = book[sheet], None, 0
            last_sku = ""
            for row in ws.iter_rows(values_only=True):
                if header is None:
                    labels = [_norm(c) for c in (row or [])]
                    if spec["order"] in labels and spec["cost"] in labels:
                        header = {label: i for i, label in enumerate(labels) if label}
                    continue

                def cell(key):
                    i = header.get(key)
                    return row[i] if i is not None and i < len(row) else None

                order_id = _order_id(cell(spec["order"]))
                sku = str(cell(spec["sku"]) or "").strip()
                cost = _number(cell(spec["cost"]))

                # The books are grouped by item and written with fill-down: only
                # the first row of a block carries the SKU, every later row for
                # the same item leaves it blank. Aug-26 has 120 rows but only 34
                # written SKUs, so discarding the blanks lost two thirds of every
                # charge column.
                if sku:
                    last_sku = sku
                elif order_id:
                    sku = last_sku

                if not (order_id and sku):
                    continue

                record = {"customer_order_id": order_id, "sku": sku,
                          "source": f"{filename}#{sheet}"}
                # A cost of 0 is an unfilled cell, not a free product.
                record["unit_cost"] = cost if (cost and cost > 0) else None
                for column, label in spec["charges"].items():
                    value = _number(cell(label))
                    # 0 is meaningful for a charge -- "no third-party label was
                    # bought" is a fact, unlike a blank unit cost.
                    record[column] = value if value is not None else None

                if record["unit_cost"] is not None or any(
                        record.get(c) for c in CHARGE_COLUMNS):
                    rows.append(record)
                    count += 1
            print(f"  {filename:32} [{sheet:16}] {count:4} rows")
        book.close()
    return rows


def main():
    parser = argparse.ArgumentParser(description="Import per-line costs from the workbooks")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])
    store_id = sb.table("stores").select("store_id").eq("active", True).execute().data[0]["store_id"]

    sheet_rows = read_books()
    print(f"  parsed {len(sheet_rows)} usable rows\n")

    def fetch_all(table, columns):
        out, offset = [], 0
        while True:
            batch = (sb.table(table).select(columns).eq("store_id", store_id)
                     .range(offset, offset + 499).execute().data)
            out += batch
            if len(batch) < 500:
                return out
            offset += 500

    orders = fetch_all("orders", "purchase_order_id,customer_order_id")
    lines = fetch_all("order_lines", "purchase_order_id,line_number,sku,status")

    customer_order = {o["purchase_order_id"]: o["customer_order_id"] for o in orders}
    db_skus = {(l["sku"] or "").strip() for l in lines}
    by_lower = {s.lower(): s for s in db_skus}

    def resolve(sku):
        for variant in sku_variants(sku):
            if variant in db_skus:
                return variant
            if variant.lower() in by_lower:
                return by_lower[variant.lower()]
        return None

    exact = {}
    per_sku = collections.defaultdict(list)
    for row in sheet_rows:
        target = resolve(row["sku"])
        if not target:
            continue
        exact[(row["customer_order_id"], target)] = row
        if row.get("unit_cost"):
            per_sku[target].append(row["unit_cost"])

    realized = [l for l in lines
                if (l["status"] or "").strip().lower() in ("shipped", "delivered")]

    out, tally = [], collections.Counter()
    for line in realized:
        sku = (line["sku"] or "").strip()
        key = (customer_order.get(line["purchase_order_id"]), sku)
        source_id = f"{line['purchase_order_id']}:{line['line_number']}"
        match = exact.get(key)
        if match and match.get("unit_cost"):
            basis, cost, source = "exact", match["unit_cost"], match["source"]
        elif per_sku.get(sku):
            costs = per_sku[sku]
            # Every row agreeing means the cost is not actually in doubt.
            basis = "sku_single" if len(set(costs)) == 1 else "sku_imputed"
            cost = statistics.median(costs)
            source = f"derived from {len(costs)} row(s) for {sku}"
        else:
            tally["no_data"] += 1
            continue
        tally[basis] += 1

        record = {"store_id": store_id, "source_id": source_id, "sku": sku,
                  "unit_cost": round(cost, 4), "basis": basis, "source": source}
        # Charges are only ever taken from this line's own row. Unlike unit cost
        # they are not a property of the SKU -- one order ships third-party and
        # the next through Walmart -- so imputing them across lines would invent
        # costs that were never incurred.
        for column in CHARGE_COLUMNS:
            value = match.get(column) if match else None
            if value is not None:
                record[column] = round(value, 4)
                tally[f"has_{column}"] += 1
        out.append(record)

    print(f"  realized lines            : {len(realized)}")
    for basis in ("exact", "sku_single", "sku_imputed", "no_data"):
        print(f"    {basis:<12} {tally[basis]:>4}")
    print("  charges captured:")
    for column in CHARGE_COLUMNS:
        total = sum(r.get(column) or 0 for r in out)
        print(f"    {column:<20} {tally['has_' + column]:>4} lines   {total:>10,.2f}")

    if args.dry_run:
        print("\n  dry run -- nothing written")
        return

    for i in range(0, len(out), 200):
        sb.table("order_line_costs").upsert(out[i:i + 200]).execute()
    print(f"\n  wrote {len(out)} order_line_costs rows")


if __name__ == "__main__":
    main()
